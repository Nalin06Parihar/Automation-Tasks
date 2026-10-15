from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _rgb_to_hex(color_value: Any) -> str | None:
    if color_value is None:
        return None
    if getattr(color_value, "type", None) != MSO_COLOR_TYPE.RGB:
        return None

    rgb = getattr(color_value, "rgb", None)
    if rgb is None:
        return None
    try:
        return "#{:02X}{:02X}{:02X}".format(rgb[0], rgb[1], rgb[2])
    except (TypeError, ValueError, IndexError):
        return None


def _font_value(font: Any, key: str) -> Any:
    if font is None:
        return None
    value = getattr(font, key, None)
    if value is None:
        return None
    if key == "size":
        return round(value.pt, 2) if value is not None else None
    return bool(value) if key in {"bold", "italic"} else value


def _collect_text_properties(shape: Any) -> dict[str, Any]:
    properties: dict[str, Any] = {}
    text_frame = getattr(shape, "text_frame", None)
    if text_frame is None:
        return properties

    properties["text"] = shape.text

    first_run = None
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text:
                first_run = run
                break
        if first_run is not None:
            break

    if first_run is not None:
        properties["font_size"] = _font_value(first_run.font, "size")
        properties["font_color"] = _rgb_to_hex(first_run.font.color)
        properties["bold"] = _font_value(first_run.font, "bold")
        properties["italic"] = _font_value(first_run.font, "italic")

    return properties


def _collect_shape_properties(shape: Any) -> dict[str, Any]:
    properties: dict[str, Any] = {}

    fill = getattr(shape, "fill", None)
    if fill is not None:
        fill_color = getattr(fill, "fore_color", None)
        if fill_color is not None:
            properties["fill_color"] = _rgb_to_hex(fill_color)

    line = getattr(shape, "line", None)
    if line is not None:
        line_color = getattr(line, "color", None)
        if line_color is not None:
            properties["line_color"] = _rgb_to_hex(line_color)

    return properties


def _safe_shape_type(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return "text"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "unsupported"
    return "shape"


def _display_name(shape: Any, properties: dict[str, Any]) -> str:
    text = " ".join(str(properties.get("text", "")).split())
    if text:
        preview = text[:48]
        if len(text) > 48:
            preview += "..."
        return f"Text: {preview}"

    try:
        return shape.auto_shape_type.name.replace("_", " ").title()
    except (AttributeError, ValueError):
        pass

    name = str(getattr(shape, "name", ""))
    if re.fullmatch(r"(shape|text|textbox)\s*\d+", name, re.IGNORECASE):
        return "Empty text box" if getattr(shape, "has_text_frame", False) else "PowerPoint shape"
    return name or "PowerPoint element"


def inspect_shape(shape: Any, slide_number: int) -> dict[str, Any] | None:
    if shape is None:
        return None

    shape_type = _safe_shape_type(shape)
    if shape_type not in {"text", "shape"}:
        return None

    properties: dict[str, Any] = {}
    if shape_type == "text":
        properties.update(_collect_text_properties(shape))
    if shape_type == "shape":
        properties.update(_collect_shape_properties(shape))
    return {
        "id": int(shape.shape_id),
        "name": getattr(shape, "name", f"Shape {shape.shape_id}"),
        "display_name": _display_name(shape, properties),
        "type": shape_type,
        "slide": slide_number,
        "left": int(getattr(shape, "left", 0)),
        "top": int(getattr(shape, "top", 0)),
        "width": int(getattr(shape, "width", 0)),
        "height": int(getattr(shape, "height", 0)),
        "properties": properties,
    }


def inspect_presentation(file_bytes: bytes, template_name: str = "uploaded_template.pptx") -> dict[str, Any]:
    presentation = Presentation(BytesIO(file_bytes))
    template: dict[str, Any] = {"template": template_name, "slides": []}

    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_entry = {"slide": slide_index, "shapes": []}
        display_name_counts: dict[str, int] = {}
        for shape in slide.shapes:
            shape_metadata = inspect_shape(shape, slide_index)
            if shape_metadata is not None:
                display_name = shape_metadata["display_name"]
                display_name_counts[display_name] = display_name_counts.get(display_name, 0) + 1
                if display_name_counts[display_name] > 1:
                    shape_metadata["display_name"] = f"{display_name} ({display_name_counts[display_name]})"
                slide_entry["shapes"].append(shape_metadata)
        template["slides"].append(slide_entry)

    return template
