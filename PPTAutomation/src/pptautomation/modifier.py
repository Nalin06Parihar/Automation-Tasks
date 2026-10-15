from __future__ import annotations

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.util import Pt


def _hex_to_rgb(hex_color: str | None) -> RGBColor | None:
    if not hex_color:
        return None

    value = hex_color.strip()
    if value.startswith("#"):
        value = value[1:]
    if len(value) == 3:
        value = "".join(ch * 2 for ch in value)
    if len(value) != 6:
        return None

    try:
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
    except ValueError:
        return None


def _set_shape_color(shape: Any, color_name: str, color_value: Any) -> None:
    if color_value in (None, ""):
        return

    rgb_color = _hex_to_rgb(str(color_value))
    if rgb_color is None:
        raise ValueError(f"Invalid color value: {color_value}")

    if color_name == "fill_color":
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb_color
    elif color_name == "line_color":
        if shape.line is not None:
            shape.line.color.rgb = rgb_color


def _apply_text_updates(shape: Any, updates: dict[str, Any]) -> None:
    if not getattr(shape, "has_text_frame", False):
        return

    text_frame = shape.text_frame
    if "text" in updates:
        text_frame.clear()
        paragraph = text_frame.paragraphs[0]
        paragraph.text = str(updates["text"])

    if "font_size" in updates:
        size = updates["font_size"]
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(float(size))

    if "font_color" in updates:
        rgb_color = _hex_to_rgb(str(updates["font_color"]))
        if rgb_color is not None:
            for paragraph in text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = rgb_color

    if "bold" in updates:
        value = bool(updates["bold"])
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.bold = value

    if "italic" in updates:
        value = bool(updates["italic"])
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.italic = value


def _apply_shape_updates(shape: Any, updates: dict[str, Any]) -> None:
    for key in ("fill_color", "line_color"):
        if key in updates:
            _set_shape_color(shape, key, updates[key])

    if getattr(shape, "has_text_frame", False):
        _apply_text_updates(shape, updates)


def apply_changes_to_presentation(
    original_presentation_bytes: bytes,
    template: dict[str, Any],
    input_values: dict[str, Any],
) -> bytes:
    presentation = Presentation(BytesIO(original_presentation_bytes))

    for slide_entry in template.get("slides", []):
        slide_number = int(slide_entry.get("slide", 0))
        if slide_number <= 0 or slide_number > len(presentation.slides):
            raise ValueError(f"Slide {slide_number} does not exist in the uploaded presentation.")

        slide = presentation.slides[slide_number - 1]
        slide_key = f"slide_{slide_number}"
        changes = input_values.get(slide_key, {})
        if not isinstance(changes, dict):
            continue

        for shape_entry in slide_entry.get("shapes", []):
            shape_name = shape_entry.get("name")
            if not shape_name:
                continue

            if isinstance(changes.get(shape_name), dict):
                updates = changes[shape_name]
                shape = None
                for candidate in slide.shapes:
                    if candidate.shape_id == int(shape_entry.get("id", -1)):
                        shape = candidate
                        break
                if shape is None:
                    for candidate in slide.shapes:
                        if candidate.name == shape_name:
                            shape = candidate
                            break
                if shape is None:
                    raise ValueError(f"Shape '{shape_name}' could not be found on Slide {slide_number}.")

                _apply_shape_updates(shape, updates)

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
